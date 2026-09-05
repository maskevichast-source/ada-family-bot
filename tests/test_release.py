import asyncio
import datetime as dt
import importlib
from unittest.mock import AsyncMock
import pytest
from services import reminders as r, debts, sheets, state
from services.timezone import ASTANA_TZ, parse_ru_relative_datetime, parse_flexible_datetime
from services.money import parse_amount
from services.weather import format_forecast
from services.telegram_safe import chunks, safe_send_message
from services.pending_receipts import set_pending, pop_pending, ack_pending, has_pending

NOW = dt.datetime(2026, 9, 5, 10, 0, tzinfo=ASTANA_TZ)

@pytest.mark.parametrize("raw,expected", [
    ("184400,50",184400.5), ("184 400",184400), ("184,400",184400),
    ("1.234,56",1234.56), ("1,234.56",1234.56), ("-500",-500), ("0",0),
    ("дал Саше 10 тыс",10000), ("2,5 тысячи",2500), ("250.0",250),
    ("1230",1230), ("Кола 500, чек 123",500), (float("nan"),0), ("текст",0),
])
def test_money(raw,expected):
    assert parse_amount(raw) == expected

@pytest.mark.parametrize("text,expected",[
    ("сегодня в 12.00", "2026-09-05 12:00"),
    ("завтра в 9 утра", "2026-09-06 09:00"),
    ("послезавтра во столько-то", None),
    ("через 15 минут", "2026-09-05 10:15"),
    ("через два часа", "2026-09-05 12:00"),
    ("через полчаса", "2026-09-05 10:30"),
    ("сегодня в 9 утра", None),
    ("в 09:00", "2026-09-06 09:00"),
    ("06.10.2026 в 12.00", "2026-10-06 12:00"),
    ("31.02.2027 в 12:00", None),
    ("завтра в 25:00", None),
    ("завтра в 12:70", None),
    ("в девять вечера", "2026-09-05 21:00"),
    ("в понедельник в 10", "2026-09-07 10:00"),
    ("завтра в 9:00 и 18:00", None),
    ("10 октября в 12:00", None),
])
def test_dates(text,expected):
    actual = parse_ru_relative_datetime(text,NOW)
    assert (actual.strftime("%Y-%m-%d %H:%M") if actual else None) == expected

def test_iso_timezone():
    assert parse_flexible_datetime("2026-09-05T06:00:00Z").hour == 11

@pytest.mark.parametrize("text,target",[
    ("наопмни диане завтра в 12.00 написать Салтанат","Диана"),
    ("Напомни нам обоим в 9", "Семья"),
    ("Напомни Диане и Владу в 9", "Семья"),
    ("Напомни мне взять у Дианы крем в 9","Влад"),
    ("Поставь напоминание для Дианы в 9","Диана"),
    ("Напомни Владу в 21:00", "Влад"),
])
def test_target(text,target):
    assert r.is_add(text)
    assert r.extract_target(text,"Влад") == target

def test_task_and_recurrence():
    text = "наопмни диане каждый день в 12.00 написать Салтанат"
    assert r.extract_task(text) == "написать салтанат"
    assert r.recurrence_from_text(text) == "daily"

def test_month_end():
    first = dt.datetime(2026,1,31,9,tzinfo=ASTANA_TZ)
    second = r.next_occurrence(first,"monthly",first,31)
    assert second.day == 28
    third = r.next_occurrence(second,"monthly",second,31)
    assert third.day == 31 and third.month == 3

def test_leap_year():
    first = dt.datetime(2028,1,31,9,tzinfo=ASTANA_TZ)
    assert r.next_occurrence(first,"monthly",first,31).day == 29

def test_skip_overdue_daily():
    assert r.next_occurrence(NOW,"daily",NOW+dt.timedelta(days=12,hours=2)) == NOW+dt.timedelta(days=13)

def test_schema_migration_and_idempotency(db,monkeypatch):
    from conftest import Sheet
    db.sheets["Reminders"] = Sheet("Reminders",[r.HEADERS[:7]],7)
    monkeypatch.setattr(r,"now_astana",lambda:NOW)
    r.add("Владислав и Диана",NOW+dt.timedelta(hours=2),"витамины",reminder_id="REM_TEST")
    r.add("Владислав и Диана",NOW+dt.timedelta(hours=2),"витамины",reminder_id="REM_TEST")
    assert len(r.pending()) == 1
    assert r.pending()[0]["target_user"] == "Семья"
    assert db.worksheet("Reminders").row_values(1) == r.HEADERS

def seed_reminder(db,target="Диана",repeat="once"):
    ws = r.worksheet()
    row = ["REM_X",NOW.isoformat(),target,NOW.strftime("%Y-%m-%d %H:%M:%S"),"Тест","pending",repeat,"Влад",5,"{}"]
    ws.append_row(row)
    return ws

def test_failed_delivery_stays_pending(db,monkeypatch):
    import config
    monkeypatch.setattr(config,"DIANA_TELEGRAM_ID","22")
    monkeypatch.setattr(config,"FAMILY_CHAT_ID",-100)
    seed_reminder(db)
    bot = type("Bot",(),{"send_message":AsyncMock(side_effect=RuntimeError("blocked"))})()
    asyncio.run(r.deliver_due(bot,NOW))
    assert len(r.pending()) == 1

def test_success_delivered_to_family_with_diana_address(db,monkeypatch):
    import config
    monkeypatch.setattr(config,"DIANA_TELEGRAM_ID","22")
    monkeypatch.setattr(config,"FAMILY_CHAT_ID",-100)
    seed_reminder(db)
    bot = type("Bot",(),{"send_message":AsyncMock()})()
    asyncio.run(r.deliver_due(bot,NOW))
    assert bot.send_message.call_args.kwargs["chat_id"] == -100
    assert "Диана" in bot.send_message.call_args.kwargs["text"]
    assert not r.pending()

def test_both_single_group_message_and_retry(db,monkeypatch):
    import config
    monkeypatch.setattr(config,"DIANA_TELEGRAM_ID","22")
    monkeypatch.setattr(config,"FAMILY_CHAT_ID",-100)
    monkeypatch.setattr(config,"VLAD_TELEGRAM_ID","11")
    seed_reminder(db,"Владислав и Диана")
    calls = []
    async def send(**kw):
        calls.append(kw["chat_id"])
        if len(calls) == 1:
            raise RuntimeError("offline")
    bot = type("Bot",(),{})()
    bot.send_message = send
    asyncio.run(r.deliver_due(bot,NOW))
    asyncio.run(r.deliver_due(bot,NOW))
    assert calls == [-100,-100]
    assert not r.pending()

def test_cancel_by_id_stable_after_other_rows(db):
    seed_reminder(db)
    r.update_by_id("REM_X",{"status":"cancelled"})
    assert r.pending() == []

def test_deletion_search_no_financial_rows(db):
    items=[{"reminder_id":"REM_1","target_user":"Диана","text":"выпить витамины","remind_at":"2026-10-01"},
           {"reminder_id":"REM_2","target_user":"Влад","text":"кешбэк","remind_at":"2026-10-01"}]
    assert [x["reminder_id"] for x in r.find_matches(items,"удали напоминание Диане про витамины","Влад")] == ["REM_1"]

def test_debt_lend_repay_partial_full(db):
    p=debts.parse_local("Дал Саше в долг 10 000","Влад")
    saved=debts.record(p,"EVENT_1")
    debts.record(p,"EVENT_1")
    assert len(debts.events()) == 1
    assert debts.balances()[0]["balance"] == 10000
    pay=debts.parse_local("Саша вернул мне 2 500 в счет долга","Влад")
    debts.record(pay,"EVENT_2")
    assert debts.balances()[0]["balance"] == 7500
    debts.record(dict(pay,amount=7500),"EVENT_3")
    assert debts.balances()[0]["balance"] == 0
    assert len(db.worksheet("Transactions").data) == 1

def test_borrow_and_return(db):
    p=debts.parse_local("Взял у Саши в долг 5000","Влад")
    debts.record(p,"1")
    repayment=debts.parse_local("Вернул Саше долг 1000","Влад")
    assert repayment["direction"]=="borrowed"
    debts.record(repayment,"2")
    assert debts.balances()[0]["balance"]==4000

def test_ambiguous_loan_rejected(db):
    with pytest.raises(ValueError):
        debts.parse_local("Одолжил Саше 1000","Влад")

def test_multiple_debts_need_id(db):
    for n in range(2):
        debts.record(debts.parse_local("Дал Саше в долг 1000","Влад"),str(n))
    with pytest.raises(ValueError):
        debts.parse_local("Саша вернул мне 500 долга","Влад")
    did=debts.balances()[0]["debt_id"]
    p=debts.parse_local(f"Погаси долг {did} на 500","Влад")
    assert p["amount"]==500 and p["debt_id"]==did

def test_overpayment_and_wrong_owner(db):
    debt=debts.record(debts.parse_local("Дал Саше в долг 1000","Влад"),"1")
    with pytest.raises(ValueError):
        debts.record(dict(debt,event_type="repay",amount=1100),"2")
    with pytest.raises(ValueError):
        debts.record(dict(debt,event_type="repay",amount=100,owner="Диана"),"3")

def test_due_date_not_counted_as_amount(db):
    p=debts.parse_local("Дал Саше в долг 1000 до 20.10.2026","Влад")
    assert p["amount"]==1000 and p["due_date"]=="2026-10-20"

def test_other_currency_rejected(db):
    with pytest.raises(ValueError):
        debts.parse_local("Дал Саше в долг 100 долларов","Влад")

def test_transactions_numeric_positive_idempotent(db):
    tx={"amount":"184400,50","transaction_id":"TX1","user":"Диана","type":"income"}
    sheets.append_transaction(tx)
    sheets.append_transaction(tx)
    row=db.worksheet("Transactions").data[-1]
    assert row[4] == 184400.5 and row[3]=="ДОХОД"
    assert len(db.worksheet("Transactions").data)==2
    with pytest.raises(ValueError): sheets.append_transaction({"amount":-100})

def test_sheets_failure_not_create_duplicate(db,monkeypatch):
    def fail(name): raise RuntimeError("403")
    monkeypatch.setattr(db,"worksheet",fail)
    with pytest.raises(RuntimeError): sheets._get_or_create_worksheet("X",["id"])
    assert "X" not in db.sheets

def test_no_amount_fallback_deletion(db):
    sheets.append_transaction({"transaction_id":"TX1","amount":500,"user_comment":"кафе"})
    sheets.append_transaction({"transaction_id":"TX2","amount":500,"user_comment":"продукты"})
    assert sheets.delete_record_by_keyword("Transactions","500") is None
    assert sheets.delete_record_by_keyword("Transactions","") is None
    with pytest.raises(ValueError): sheets.delete_record_by_keyword("Reminders","REM_X")

def test_blank_rows_not_selected_as_last(db):
    sheets.append_transaction({"transaction_id":"TX1","amount":500})
    db.worksheet("Transactions").data.extend([[""]*16 for _ in range(3)])
    rec=sheets.delete_record_by_keyword("Transactions","последняя")
    assert rec["transaction_id"]=="TX1"

def test_subscription_day31_feb_idempotence(db):
    ws=sheets._get_or_create_subscriptions_sheet()
    ws.append_row(["SUB_X","Test",1000,"Kaspi",31,"2026-01-31","active",""])
    now=dt.datetime(2026,2,28,10,tzinfo=ASTANA_TZ)
    assert sheets.process_due_subscriptions(now)==["Test"]
    assert sheets.process_due_subscriptions(now)==[]
    assert len(db.worksheet("Transactions").data)==2

def test_pending_isolation_ack_restart(monkeypatch):
    set_pending("-1:11",[{"amount":100,"transaction_id":"x"}],"Влад")
    set_pending("-1:22",[{"amount":200,"transaction_id":"y"}],"Диана")
    assert pop_pending("-1:11")[1]=="Влад"
    import services.pending_receipts as pr
    importlib.reload(pr)
    assert pr.has_pending("-1:11")
    pr.ack_pending("-1:11")
    assert pr.has_pending("-1:22")

def test_pending_multiple_receipts_not_overwritten():
    set_pending("1:1",[{"amount":100}],"Влад")
    set_pending("1:1",[{"amount":200}],"Влад")
    assert len(pop_pending("1:1")[0])==2

def test_weather_missing_data_and_no_past_slots():
    data={"daily":{"time":["2026-09-05"],"temperature_2m_min":[0],"temperature_2m_max":[12]},
          "hourly":{"time":["2026-09-05T09:00","2026-09-05T19:00"],"temperature_2m":[5,10]}}
    text=format_forecast(data,"today",NOW.replace(hour=22))
    assert "09:00" not in text and "19:00" not in text
    assert "None%" not in text and "нет данных" in text
    assert "\n\n" in text

def test_week_has_separate_blocks():
    data={"daily":{"time":["2026-09-05","2026-09-06"],"temperature_2m_min":[0,None]}}
    text=format_forecast(data,"week",NOW)
    assert "Сегодня" in text and "Завтра" in text
    assert "\n\n" in text

def test_long_messages_chunk():
    text="abc\n"*3000
    parts=list(chunks(text))
    assert len(parts)>1 and max(map(len,parts))<=3500

def test_telegram_bad_request_fallback():
    from aiogram.exceptions import TelegramBadRequest
    err=TelegramBadRequest(method=None,message="can't parse entities")
    bot=type("Bot",(),{"send_message":AsyncMock(side_effect=[err,True])})()
    assert asyncio.run(safe_send_message(bot,22,"bad_*")) is True
    assert bot.send_message.call_args.kwargs["parse_mode"] is None

class Message:
    def __init__(self, uid=11, mid=1):
        from types import SimpleNamespace
        self.chat=SimpleNamespace(id=-100)
        self.from_user=SimpleNamespace(id=uid)
        self.message_id=mid
        self.answer=AsyncMock()

def test_reminder_dialogue_time_followup(db,monkeypatch):
    monkeypatch.setattr(r,"now_astana",lambda:NOW)
    monkeypatch.setattr(r,"parse_ru_relative_datetime",lambda text:parse_ru_relative_datetime(text,NOW))
    msg=Message()
    assert not asyncio.run(r.handle(msg,"наопмни диане написать Салтанат","Влад"))  # falls through to conversational model
    msg.message_id=2
    assert asyncio.run(r.handle(msg,"завтра в 12.00","Влад"))
    assert r.pending()[0]["target_user"]=="Диана"
    assert r.pending()[0]["remind_at"]=="2026-09-06 12:00:00"

def test_reminder_delete_requires_confirmation(db):
    seed_reminder(db)
    msg=Message()
    asyncio.run(r.handle(msg,"удали напоминание REM_X","Влад"))
    assert len(r.pending())==1
    asyncio.run(r.handle(msg,"да","Влад"))
    assert not r.pending()
    assert db.worksheet("Reminders").data[1][5]=="cancelled"

def test_debt_requires_confirmation(db):
    msg=Message()
    asyncio.run(debts.handle(msg,"Дал Саше в долг 1000","Влад"))
    assert debts.balances()==[]
    asyncio.run(debts.handle(msg,"да","Влад"))
    assert debts.balances()[0]["balance"]==1000

def test_stranger_alias_not_authorized(monkeypatch):
    import config
    monkeypatch.setattr(config,"VLAD_TELEGRAM_ID","11")
    assert config.get_authorized_user_name(99,"Влад") is None


def test_atomic_split_request(db):
    sheets.append_transaction({"transaction_id":"TRX1","amount":1000})
    db.worksheet("Transactions").id = 123
    requests=[]
    db.batch_update=lambda payload:requests.append(payload)
    assert sheets.split_last_transaction_by_amount(1000,400,"Еда и продукты","а",600,"Еда и продукты","б")
    batch=requests[0]["requests"]
    assert len(batch)==2 and "insertDimension" in batch[0] and "updateCells" in batch[1]
    assert batch[1]["updateCells"]["start"]["rowIndex"]==1
    assert sheets.split_last_transaction_by_amount(1000,400,"Еда и продукты","а",700,"Еда и продукты","б") is False
    assert len(requests)==1

def test_multiple_debt_amounts_rejected(db):
    with pytest.raises(ValueError):
        debts.parse_local("Дал Саше 1000 и Пете 2000 в долг","Влад")


def test_excel_export_keeps_actual_fields_and_formula_is_text(db):
    import io
    import openpyxl
    from services.reports import generate_excel_export
    from services.timezone import now_astana
    now=now_astana()
    sheets.append_transaction({"transaction_id":"EXPORT1","amount":100,"source":"Кошелек",
        "resource":"Наличные","funds_type":"Кредитные","merchant":"=HYPERLINK(\"x\")",
        "ai_comment":"примечание","user_comment":"=1+1"})
    content=generate_excel_export(now.year,now.month)
    book=openpyxl.load_workbook(io.BytesIO(content))
    row=list(book["Выписка"].values)[1]
    assert row[8]=="Кредитные" and row[9]=="Наличные"
    assert row[12]=='=HYPERLINK("x")' and row[15]=="примечание"
    assert book["Выписка"]["M2"].data_type=="s"
    assert "Журнал долгов" in book.sheetnames and "Остатки долгов" in book.sheetnames

def test_foreign_currency_cannot_silently_mix(db):
    with pytest.raises(ValueError):
        sheets.append_transaction({"amount":100,"currency":"USD"})


def test_reminder_migration_refuses_unrelated_columns(db):
    from conftest import Sheet
    db.sheets["Reminders"]=Sheet("Reminders",[r.HEADERS[:7]+["custom_data"]],8)
    with pytest.raises(RuntimeError): r.worksheet()
    assert db.sheets["Reminders"].data[0][7]=="custom_data"
