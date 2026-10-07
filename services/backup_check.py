"""Ежемесячная проверка, что резервная копия таблицы не пустышка.

Бот выгружает таблицу в .xlsx (так же, как ночной бэкап), открывает файл как чужой, находит лист Transactions и
сравнивает число записей с живой таблицей. Файл никуда не сохраняется. Отчёт уходит Владу в личку."""
import io

REQUIRED_SHEETS = ("Transactions",)
TOLERANCE_ROWS = 5            # пока шла выгрузка, в таблицу могли дописать пару трат


def _count_rows(rows) -> int:
    """Число строк данных (без заголовка), где заполнена первая колонка (transaction_id)."""
    total = 0
    for i, row in enumerate(rows):
        if i == 0:
            continue
        first = row[0] if row else None
        if first not in (None, ""):
            total += 1
    return total


def compare(exported_rows: int, live_rows: int, sheet_names: list[str]) -> dict:
    missing = [name for name in REQUIRED_SHEETS if name not in sheet_names]
    if missing:
        return {"ok": False, "reason": f"в файле нет листа {', '.join(missing)}", "exported": exported_rows, "live": live_rows}
    if live_rows - exported_rows > TOLERANCE_ROWS:
        return {"ok": False, "reason": f"в копии {exported_rows} записей, а в таблице {live_rows}: копия неполная",
                "exported": exported_rows, "live": live_rows}
    if exported_rows - live_rows > TOLERANCE_ROWS:
        return {"ok": False, "reason": f"в копии {exported_rows} записей, а в таблице {live_rows}: числа не сходятся",
                "exported": exported_rows, "live": live_rows}
    return {"ok": True, "exported": exported_rows, "live": live_rows, "sheets": len(sheet_names)}


def verify_backup() -> dict:
    """{"ok": True, "exported", "live", "sheets"} или {"ok": False, "reason"}."""
    try:
        from gspread.utils import ExportFormat
        from openpyxl import load_workbook
        from services import sheets
        db = sheets.get_db()
        data = sheets.get_client().export(db.id, ExportFormat.EXCEL)
        if not data or bytes(data[:2]) != b"PK":
            return {"ok": False, "reason": "Google вернул пустой или повреждённый файл"}
        workbook = load_workbook(io.BytesIO(bytes(data)), read_only=True, data_only=True)
        names = list(workbook.sheetnames)
        exported = _count_rows(workbook["Transactions"].iter_rows(values_only=True)) if "Transactions" in names else 0
        live = _count_rows(db.worksheet("Transactions").get_all_values())
        return compare(exported, live, names)
    except Exception as error:  # noqa: BLE001 - причина показывается Владу
        return {"ok": False, "reason": f"{type(error).__name__}: {str(error)[:200]}"}


def report_text(result: dict) -> str:
    if result.get("ok"):
        return (f"Проверка бэкапа: всё в порядке. Копию таблицы удалось открыть, листов {result['sheets']}, "
                f"записей в Transactions {result['exported']} (в живой таблице {result['live']}).")
    return f"Проверка бэкапа: ПРОБЛЕМА. {result.get('reason')}. Сделай копию вручную командой /backup и проверь таблицу."
