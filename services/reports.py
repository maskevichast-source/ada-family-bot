"""Генерация PDF-отчётов и выгрузки Excel (.xlsx) прямо в Telegram."""

import io
import re
import textwrap
import datetime
from collections import defaultdict
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from services.sheets import get_transactions_for_period, get_category_limits
from services.categories import TYPE_EXPENSE, TYPE_INCOME
from services.timezone import now_astana
from services.money import parse_amount
from services.analytics import analyze_budget_leaks
from services.plot_lock import serialized_plot


_EMOJI_RE = re.compile(
    "["
    "\U0001F300-\U0001FAFF"
    "\U00002600-\U000027BF"
    "\U0001F1E6-\U0001F1FF"
    "\U00002190-\U000021FF"
    "]+",
    flags=re.UNICODE,
)


def _wrap_pdf(text: str, width: int = 80) -> str:
    """Готовит текст для вставки в PDF через matplotlib: убирает эмодзи
    (у стандартного DejaVu Sans нет этих глифов — вместо них печатаются
    пустые квадраты, и matplotlib на каждый такой символ шлёт warning) и
    переносит длинные строки настоящими переводами строки, а не тем, что
    легко случайно оставить как литеральный текст "\\n" внутри f-строки."""
    clean = _EMOJI_RE.sub("", str(text or ""))
    clean = re.sub(r"[ \t]+", " ", clean)
    out_lines = []
    for line in clean.split("\n"):
        line = line.strip()
        if not line:
            out_lines.append("")
            continue
        out_lines.extend(textwrap.wrap(line, width=width) or [""])
    return "\n".join(out_lines)


def _format_currency(value):
    try:
        return f"{float(value):,.0f}".replace(",", " ")
    except (ValueError, TypeError):
        return str(value)


@serialized_plot
def generate_pdf_report(year: int = None, month: int = None) -> bytes:
    """Генерирует двухстраничный финансовый PDF-буклет семьи."""
    now = now_astana()
    y = year or now.year
    m = month or now.month

    start_date = datetime.date(y, m, 1).strftime("%Y-%m-%d")
    if m == 12:
        end_date = datetime.date(y + 1, 1, 1).strftime("%Y-%m-%d")
    else:
        end_date = datetime.date(y, m + 1, 1).strftime("%Y-%m-%d")

    transactions = get_transactions_for_period(start_date, end_date)
    month_name = datetime.date(y, m, 1).strftime("%B %Y").capitalize()

    expenses = [t for t in transactions if str(t.get("type", "")).strip() != TYPE_INCOME]
    incomes = [t for t in transactions if str(t.get("type", "")).strip() == TYPE_INCOME]

    total_exp = sum(parse_amount(t.get("amt", 0)) for t in expenses)
    total_inc = sum(parse_amount(t.get("amt", 0)) for t in incomes)
    balance = total_inc - total_exp

    buf = io.BytesIO()

    with PdfPages(buf) as pdf:
        # ── СТРАНИЦА 1: ДАШБОРД МЕСЯЦА ──
        fig1 = plt.figure(figsize=(8.27, 11.69), facecolor="#1a1a2e")  # А4
        fig1.suptitle(f"СЕМЕЙНЫЙ ФИНАНСОВЫЙ ОТЧЁТ\n{month_name}", fontsize=18, fontweight="bold", color="#ecf0f1", y=0.95)

        # 1. Текстовые карточки (Доход / Расход / Баланс)
        ax_cards = fig1.add_subplot(3, 1, 1)
        ax_cards.set_facecolor("#1a1a2e")
        ax_cards.axis("off")
        card_text = _wrap_pdf(
            f"Общий доход:  {_format_currency(total_inc)} KZT\n"
            f"Общий расход: {_format_currency(total_exp)} KZT\n"
            f"Чистый баланс: {_format_currency(balance)} KZT\n"
            f"Всего операций: {len(transactions)} (Расходов: {len(expenses)}, Доходов: {len(incomes)})"
        )
        ax_cards.text(0.1, 0.4, card_text, fontsize=12, color="#ecf0f1", linespacing=1.6,
                      bbox=dict(boxstyle="round,pad=1", facecolor="#27293d", edgecolor="#4ECDC4", linewidth=1.5))

        # 2. Круговая диаграмма расходов
        by_cat = defaultdict(float)
        for t in expenses:
            by_cat[str(t.get("cat", "Прочее"))] += parse_amount(t.get("amt", 0))

        ax_pie = fig1.add_subplot(3, 2, 3)
        ax_pie.set_facecolor("#1a1a2e")
        top_cats = sorted(by_cat.items(), key=lambda x: -x[1])[:6]
        other_amt = sum(v for _, v in sorted(by_cat.items(), key=lambda x: -x[1])[6:])
        if other_amt > 0:
            top_cats.append(("Прочее", other_amt))

        labels = [c[0][:14] for c in top_cats]
        vals = [c[1] for c in top_cats]
        colors = ['#FF6B6B', '#4ECDC4', '#45B7D1', '#96CEB4', '#FFEAA7', '#DDA0DD', '#85C1E9']

        if sum(vals) > 0:
            ax_pie.pie(vals, labels=labels, autopct="%1.0f%%", colors=colors[:len(vals)], textprops={"color": "#ecf0f1", "fontsize": 8})
        else:
            ax_pie.text(0.5, 0.5, "Трат нет", color="#ecf0f1", ha="center")
        ax_pie.set_title("Расходы по категориям", color="#ecf0f1", fontsize=10, pad=8)

        # 3. Need vs Want
        ax_nec = fig1.add_subplot(3, 2, 4)
        ax_nec.set_facecolor("#1a1a2e")
        by_nec = defaultdict(float)
        for t in expenses:
            by_nec[str(t.get("nec", "Want"))] += parse_amount(t.get("amt", 0))

        n_labels = list(by_nec.keys())
        n_vals = list(by_nec.values())
        n_colors = ['#2ECC71' if 'need' in str(l).lower() else '#E74C3C' for l in n_labels]
        ax_nec.bar(n_labels, n_vals, color=n_colors, width=0.4)
        ax_nec.set_title("Need vs Want", color="#ecf0f1", fontsize=10, pad=8)
        ax_nec.tick_params(colors="#ecf0f1", labelsize=8)
        ax_nec.grid(axis="y", alpha=0.15)

        pdf.savefig(fig1)
        plt.close(fig1)

        # ── СТРАНИЦА 2: УТЕЧКИ БЮДЖЕТА И ТОП ТРАТ ──
        fig2 = plt.figure(figsize=(8.27, 11.69), facecolor="#1a1a2e")
        fig2.suptitle("АНАЛИЗ ПРИВЫЧЕК И КРУПНЫЕ ТРАТЫ", fontsize=16, fontweight="bold", color="#ecf0f1", y=0.95)

        # Утечки бюджета
        ax_leaks = fig2.add_subplot(2, 1, 1)
        ax_leaks.set_facecolor("#1a1a2e")
        ax_leaks.axis("off")
        leak_data = analyze_budget_leaks(y, m)
        ax_leaks.text(0.05, 0.2, _wrap_pdf(leak_data["text"].replace("**", "").replace("_", ""), 90),
                      fontsize=10, color="#ecf0f1", linespacing=1.5,
                      bbox=dict(boxstyle="round,pad=1", facecolor="#27293d", edgecolor="#E74C3C", linewidth=1.5))

        # Топ-7 крупных покупок месяца
        ax_top = fig2.add_subplot(2, 1, 2)
        ax_top.set_facecolor("#1a1a2e")
        ax_top.axis("off")
        top_txs = sorted(expenses, key=lambda x: -parse_amount(x.get("amt", 0)))[:7]
        top_lines = ["Топ-7 крупнейших трат месяца:\n"]
        for idx, t in enumerate(top_txs, 1):
            top_lines.append(f"{idx}. {_format_currency(t.get('amt'))} KZT — {t.get('cat')} ({t.get('comm')}) | {t.get('user')}")

        ax_top.text(0.05, 0.2, "\n".join(top_lines), fontsize=10, color="#ecf0f1", linespacing=1.6,
                    bbox=dict(boxstyle="round,pad=1", facecolor="#27293d", edgecolor="#45B7D1", linewidth=1.5))

        pdf.savefig(fig2)
        plt.close(fig2)

    buf.seek(0)
    return buf.getvalue()


def generate_excel_export(year: int = None, month: int = None) -> bytes:
    """Генерирует чистую выписку Excel (.xlsx) со всеми 16 колонками и стилями."""
    now = now_astana()
    y = year or now.year
    m = month or now.month

    start_date = datetime.date(y, m, 1).strftime("%Y-%m-%d")
    if m == 12:
        end_date = datetime.date(y + 1, 1, 1).strftime("%Y-%m-%d")
    else:
        end_date = datetime.date(y, m + 1, 1).strftime("%Y-%m-%d")

    transactions = get_transactions_for_period(start_date, end_date)

    wb = openpyxl.Workbook()

    # Лист 1: Все операции
    ws1 = wb.active
    ws1.title = "Выписка"

    headers = [
        "ID", "Дата и время", "Пользователь", "Тип", "Сумма (KZT)", "Валюта",
        "Банк", "Источник", "Тип средств", "Ресурс", "Категория", "Подкатегория",
        "Продавец", "Статус", "Комментарий", "Комментарий Ады"
    ]
    ws1.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for t in transactions:
        merchant_val = str(t.get("merchant", ""))
        row_idx = ws1.max_row + 1
        ws1.append([
            str(t.get("transaction_id", "")),
            str(t.get("date", "")),
            str(t.get("user", "")),
            str(t.get("type", "")),
            parse_amount(t.get("amt", 0)),
            str(t.get("curr", "KZT")),
            str(t.get("bank", "")),
            str(t.get("source", "")),
            str(t.get("funds_type", "Собственные")),
            str(t.get("resource", "Карта")),
            str(t.get("cat", "")),
            str(t.get("subcat", "")),
            merchant_val,
            str(t.get("nec", "")),
            str(t.get("comm", "")),
            str(t.get("ai_comm", "")),
        ])
        if merchant_val.startswith("="):
            # Название продавца — это ДАННЫЕ, а не формула, даже если
            # начинается с "=" (например, кто-то назвал товар с таким
            # символом). Без этого Excel попытается это вычислить.
            # ВАЖНО: сначала meняем data_type, потом НЕ трогаем .value —
            # повторное присвоение .value заново запускает автоопределение
            # типа и откатывает обратно на "формулу".
            cell = ws1.cell(row=row_idx, column=13)
            cell.data_type = "s"

    # Автоподбор ширины столбцов
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 40)

    # Лист 2: Журнал долгов (сырая лента событий — открытие/погашение)
    try:
        from services import debts as debts_svc
        debt_events = debts_svc.events()
    except Exception:
        debt_events = []
    ws2 = wb.create_sheet("Журнал долгов")
    ws2.append(["Дата", "Кому/от кого", "Направление", "Тип", "Сумма", "Срок", "Заметка"])
    for e in debt_events:
        ws2.append([
            str(e.get("date", "")), str(e.get("counterparty", "")),
            "Дал в долг" if e.get("direction") == "lent" else "Взял в долг",
            "Открытие" if e.get("event_type") == "open" else "Погашение",
            parse_amount(e.get("amount", 0)), str(e.get("due_date", "")), str(e.get("note", "")),
        ])

    # Лист 3: Остатки долгов (текущий баланс по каждому)
    try:
        debt_balances = debts_svc.balances(debt_events)
    except Exception:
        debt_balances = []
    ws3 = wb.create_sheet("Остатки долгов")
    ws3.append(["Кому/от кого", "Направление", "Остаток", "Срок"])
    for b in debt_balances:
        if b.get("balance", 0) <= 0:
            continue
        ws3.append([
            str(b.get("counterparty", "")),
            "Нам должны" if b.get("direction") == "lent" else "Мы должны",
            b.get("balance", 0), str(b.get("due_date", "")),
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

